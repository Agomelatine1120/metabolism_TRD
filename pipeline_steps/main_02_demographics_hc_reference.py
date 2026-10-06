"""Auto-extracted pipeline step: Main 02 — demographics and HC-referenced deviations.

Do not edit generated copies independently of the source notebooks.
"""


# Source cell 11
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def style_excel_workbook(path):
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.print_title_rows = "1:1"
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 30
        for column_index in range(1, worksheet.max_column + 1):
            values = [worksheet.cell(row, column_index).value for row in range(1, min(worksheet.max_row, 150) + 1)]
            width = max((len(str(value)) for value in values if value is not None), default=10) + 2
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(width, 12), 34)
            header = str(worksheet.cell(1, column_index).value)
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row, column_index)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if "p_value" in header:
                    cell.number_format = "0.0000E+00"
                elif header in {"statistic", "df", "df_1", "df_2", "TRD1_vs_TRD2_t", "TRD1_vs_TRD2_df"}:
                    cell.number_format = "0.0000"
        worksheet.print_area = worksheet.dimensions
    workbook.save(path)


FOUR_GROUPS = ["HC", "nTRD", "TRD1", "TRD2"]
PATIENT_GROUPS = ["nTRD", "TRD1", "TRD2"]


def continuous_values(column, groups):
    return [
        clinical.loc[clinical["analysis_cohort"] == group, column].dropna().to_numpy(float)
        for group in groups
    ]


def format_mean_sd(values):
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]
    if not len(values):
        return "NA"
    return f"{values.mean():.2f} ({values.std(ddof=1):.2f}); n={len(values)}"


def anova_row(variable, column, groups):
    arrays = continuous_values(column, groups)
    result = stats.f_oneway(*arrays)
    row = {"variable": variable}
    row.update({group: format_mean_sd(values) for group, values in zip(groups, arrays)})
    row.update({
        "test": "one-way ANOVA",
        "statistic": float(result.statistic),
        "df_1": len(groups) - 1,
        "df_2": int(sum(len(values) for values in arrays) - len(groups)),
        "p_value": float(result.pvalue),
    })
    return row


def welch_p(column, group_1, group_2):
    values_1, values_2 = continuous_values(column, [group_1, group_2])
    return float(stats.ttest_ind(values_1, values_2, equal_var=False).pvalue)


def demographic_row(variable, column):
    arrays = continuous_values(column, ["HC", "nTRD", "TRD1"])
    row = {"variable": variable}
    row.update({group: format_mean_sd(values) for group, values in zip(FOUR_GROUPS, continuous_values(column, FOUR_GROUPS))})
    row.update({
        "primary_test": "one-way ANOVA: HC, nTRD, TRD1",
        "primary_p_value": float(stats.f_oneway(*arrays).pvalue),
        "TRD1_vs_TRD2_test": "Welch t-test",
        "TRD1_vs_TRD2_p_value": welch_p(column, "TRD1", "TRD2"),
    })
    return row


demographic_rows = [
    demographic_row("Age", "age"),
    demographic_row("Education", "education"),
]
sex_table = pd.crosstab(clinical["analysis_cohort"], clinical["sex"]).reindex(
    index=FOUR_GROUPS, columns=[0.0, 1.0], fill_value=0
)
sex_primary = stats.chi2_contingency(sex_table.loc[["HC", "nTRD", "TRD1"]], correction=False)
sex_pairwise = stats.fisher_exact(sex_table.loc[["TRD1", "TRD2"]].to_numpy())
sex_row = {"variable": "Female, n (%)"}
for group in FOUR_GROUPS:
    female = int(sex_table.loc[group, 1.0])
    total = int(sex_table.loc[group].sum())
    sex_row[group] = f"{female} ({100 * female / total:.2f}%); n={total}"
sex_row.update({
    "primary_test": "Pearson chi-square: HC, nTRD, TRD1",
    "primary_p_value": float(sex_primary.pvalue),
    "TRD1_vs_TRD2_test": "Fisher exact test",
    "TRD1_vs_TRD2_p_value": float(sex_pairwise.pvalue),
})
demographic_table = pd.DataFrame([demographic_rows[0], sex_row, demographic_rows[1]])

baseline_hdrs_table = pd.DataFrame([{
    "variable": "Baseline HDRS-17",
    "nTRD": format_mean_sd(continuous_values("hdrs_pre", ["nTRD"])[0]),
    "TRD1": format_mean_sd(continuous_values("hdrs_pre", ["TRD1"])[0]),
    "TRD2": format_mean_sd(continuous_values("hdrs_pre", ["TRD2"])[0]),
    "primary_test": "Welch t-test: nTRD versus TRD1",
    "primary_p_value": welch_p("hdrs_pre", "nTRD", "TRD1"),
    "TRD1_vs_TRD2_test": "Welch t-test",
    "TRD1_vs_TRD2_p_value": welch_p("hdrs_pre", "TRD1", "TRD2"),
}])

msm_table = pd.DataFrame([{
    "variable": "MSM scores", "note": "Not required for the released analyses"
}])

pet_paired_keys = set(pet_pre["subject_key"]).intersection(set(pet_post["subject_key"]))
paired_clinical = clinical.loc[
    clinical["analysis_cohort"].eq("TRD2") & clinical["subject_key"].isin(pet_paired_keys)
].copy()
assert paired_clinical["analysis_arm"].value_counts().to_dict() == EXPECTED_PAIRED_ARMS

reduction_arrays = []
hdrs_reduction_row = {"variable": "HDRS reduction (baseline - post)"}
for arm in EXPECTED_ARMS:
    values = paired_clinical.loc[
        paired_clinical["analysis_arm"] == arm,
        "hdrs_improvement",
    ].dropna().to_numpy(float)
    reduction_arrays.append(values)
    hdrs_reduction_row[arm] = format_mean_sd(values)
reduction_test = stats.f_oneway(*reduction_arrays)
hdrs_reduction_row.update({
    "test": "ordinary one-way ANOVA",
    "statistic": float(reduction_test.statistic),
    "df_1": len(reduction_arrays) - 1,
    "df_2": int(sum(len(values) for values in reduction_arrays) - len(reduction_arrays)),
    "p_value": float(reduction_test.pvalue),
})
hdrs_reduction_table = pd.DataFrame([hdrs_reduction_row])

stimulation_rows = []
for variable, column in [
    ("Age, years", "age"),
    ("Education, years", "education"),
    ("Baseline HDRS-17", "hdrs_pre"),
    ("HDRS-17 reduction", "hdrs_improvement"),
]:
    arrays = [
        paired_clinical.loc[paired_clinical.analysis_arm.eq(arm), column].dropna().to_numpy(float)
        for arm in ["rTMS", "iTBS", "sham"]
    ]
    stimulation_rows.append({
        "variable": variable,
        "rTMS": format_mean_sd(arrays[0]),
        "iTBS": format_mean_sd(arrays[1]),
        "sham": format_mean_sd(arrays[2]),
        "test": "one-way ANOVA",
        "p_value": float(stats.f_oneway(*arrays).pvalue),
    })
arm_sex = pd.crosstab(paired_clinical.analysis_arm, paired_clinical.sex).reindex(
    index=["rTMS", "iTBS", "sham"], columns=[0.0, 1.0], fill_value=0
)
arm_sex_test = stats.chi2_contingency(arm_sex, correction=False)
stimulation_sex = {"variable": "Female, n (%)", "test": "Pearson chi-square", "p_value": float(arm_sex_test.pvalue)}
for arm in ["rTMS", "iTBS", "sham"]:
    female = int(arm_sex.loc[arm, 1.0])
    total = int(arm_sex.loc[arm].sum())
    stimulation_sex[arm] = f"{female} ({100 * female / total:.2f}%); n={total}"
stimulation_table = pd.DataFrame([
    stimulation_rows[0], stimulation_sex, stimulation_rows[1],
    stimulation_rows[2], stimulation_rows[3],
])

subject_demographic = clinical[[
    "subject_key", "analysis_cohort", "analysis_arm", "age", "sex", "education", "tiv",
    "hdrs_pre", "hdrs_post", "hdrs_improvement", "msm_t", "msm_total",
]].copy()

demographic_file = OUTPUT_ROOT / "demographic_and_clinical_summary.xlsx"
with pd.ExcelWriter(demographic_file, engine="openpyxl") as writer:
    demographic_table.to_excel(writer, sheet_name="demographic", index=False)
    baseline_hdrs_table.to_excel(writer, sheet_name="baseline_HDRS", index=False)
    msm_table.to_excel(writer, sheet_name="MSM_scores", index=False)
    hdrs_reduction_table.to_excel(writer, sheet_name="HDRS_reduction", index=False)
    stimulation_table.to_excel(writer, sheet_name="stimulation_paired", index=False)
    subject_demographic.to_excel(writer, sheet_name="subject_level", index=False)
style_excel_workbook(demographic_file)

assert len(subject_demographic) == sum(EXPECTED_COUNTS.values())
assert clinical.groupby("analysis_cohort").size().to_dict() == EXPECTED_COUNTS
# MSM is not required for this minimized analytical input.
print("Written:", demographic_file)
display(demographic_table)
display(baseline_hdrs_table)
display(msm_table)
display(hdrs_reduction_table)
display(stimulation_table)


# Source cell 13
def individual_edge_deviation(rows, reference_values, reference_covariates, pcc_reference, denominator, covariates):
    output = []
    for _, row in rows.iterrows():
        values = np.vstack([reference_values, row[ROI_COLUMNS].to_numpy(float)])
        cov = np.vstack([reference_covariates, row[covariates].to_numpy(float)])
        perturbed = partial_corr_matrix(values, cov)
        z_full = (perturbed - pcc_reference)[tri_full] / denominator
        output.append(z_full[CROSS_EDGE_IDS])
    return np.asarray(output, dtype=float)


def build_deviation_context(case_cohort, reference_covariates, output_dir):
    """Build an HC reference and subject-level deviations without a group test."""
    output_dir.mkdir(parents=True, exist_ok=True)
    hc = cohort_pet("HC")
    case = cohort_pet(case_cohort)

    for frame, label in [(hc, "HC"), (case, case_cohort)]:
        frame.dropna(subset=reference_covariates, inplace=True)
        assert len(frame) == EXPECTED_COUNTS[label], (label, len(frame))

    reference_values = hc[ROI_COLUMNS].to_numpy(float)
    reference_cov = hc[reference_covariates].to_numpy(float)
    pcc_reference = partial_corr_matrix(reference_values, reference_cov)
    denominator = (1 - pcc_reference[tri_full] ** 2) / (len(hc) - 1)
    denominator[np.abs(denominator) < 1e-12] = 1e-12

    pd.DataFrame(pcc_reference, index=ROI_COLUMNS, columns=ROI_COLUMNS).to_excel(
        output_dir / "PCC_ref_HC.xlsx"
    )
    roi_definition.to_excel(output_dir / "roi_definition_12x14.xlsx", index=False)
    manifest = pd.concat([
        hc[["subject_key", "age", "sex", "education"]].assign(cohort="HC"),
        case[["subject_key", "age", "sex", "education"]].assign(cohort=case_cohort),
    ], ignore_index=True)
    manifest.sort_values(["cohort", "subject_key"]).to_excel(
        output_dir / "sample_manifest.xlsx", index=False
    )

    case_edge = individual_edge_deviation(
        case, reference_values, reference_cov, pcc_reference, denominator, reference_covariates
    )
    np.savez_compressed(
        output_dir / f"individual_MCM_{case_cohort}_cross168.npz",
        subject_keys=case["subject_key"].to_numpy(),
        edge_names=np.asarray(CROSS_EDGE_NAMES, dtype=object),
        z_edges=case_edge,
    )

    # Formal feature-wise HC-only OLS: regional SUVR ~ age + sex.
    regional_hc_design = sm.add_constant(hc[["age", "sex"]], has_constant="add")
    regional_beta = np.linalg.pinv(regional_hc_design.to_numpy(float)) @ hc[ROI_COLUMNS].to_numpy(float)
    regional_hc_residual = hc[ROI_COLUMNS].to_numpy(float) - regional_hc_design.to_numpy(float) @ regional_beta
    reference_sd = pd.Series(
        np.std(regional_hc_residual, axis=0, ddof=1), index=ROI_COLUMNS
    ).replace(0, np.nan)
    reference_mean = hc[ROI_COLUMNS].mean(axis=0)  # retained for compatible exports
    case_design = sm.add_constant(case[["age", "sex"]], has_constant="add").reindex(
        columns=regional_hc_design.columns
    )
    case_regional = pd.DataFrame(
        (case[ROI_COLUMNS].to_numpy(float) - case_design.to_numpy(float) @ regional_beta)
        / reference_sd.to_numpy(float),
        columns=ROI_COLUMNS,
        index=case.index,
    )
    case_subjects = case_regional.copy()
    case_subjects.insert(0, "subject_key", case["subject_key"].to_numpy())
    case_subjects.to_excel(
        output_dir / f"regional_{case_cohort}_subjects.xlsx", index=False
    )

    return {
        "case": case_cohort,
        "output_dir": output_dir,
        "hc": hc,
        "case_rows": case,
        "reference_values": reference_values,
        "reference_covariates": reference_cov,
        "reference_covariate_names": reference_covariates,
        "pcc_reference": pcc_reference,
        "denominator": denominator,
        "reference_mean": reference_mean,
        "reference_sd": reference_sd,
        "regional_hc_design_columns": list(regional_hc_design.columns),
        "regional_beta": regional_beta,
        "case_edge": case_edge,
        "case_regional": case_regional,
    }


def run_trd1_vs_ntrd(reference_covariates, glm_covariates, output_dir):
    """Run the only case-control inference: TRD1 vs nTRD, reported at p<0.01."""
    context = build_deviation_context("TRD1", reference_covariates, output_dir)
    ntrd = cohort_pet("nTRD")
    ntrd.dropna(subset=reference_covariates + glm_covariates, inplace=True)
    assert len(ntrd) == EXPECTED_COUNTS["nTRD"], len(ntrd)

    # Rewrite the manifest so it explicitly records all subjects in this comparison.
    manifest = pd.concat([
        context["hc"][["subject_key", "age", "sex", "education"]].assign(cohort="HC"),
        ntrd[["subject_key", "age", "sex", "education"]].assign(cohort="nTRD"),
        context["case_rows"][["subject_key", "age", "sex", "education"]].assign(cohort="TRD1"),
    ], ignore_index=True)
    manifest.sort_values(["cohort", "subject_key"]).to_excel(
        output_dir / "sample_manifest.xlsx", index=False
    )

    ntrd_edge = individual_edge_deviation(
        ntrd,
        context["reference_values"],
        context["reference_covariates"],
        context["pcc_reference"],
        context["denominator"],
        reference_covariates,
    )
    np.savez_compressed(
        output_dir / "individual_MCM_nTRD_cross168.npz",
        subject_keys=ntrd["subject_key"].to_numpy(),
        edge_names=np.asarray(CROSS_EDGE_NAMES, dtype=object),
        z_edges=ntrd_edge,
    )

    case = context["case_rows"]
    case_edge = context["case_edge"]
    covariate_table = clinical.set_index("subject_key")[["age", "sex", "education"]]
    all_keys = list(case["subject_key"]) + list(ntrd["subject_key"])
    group = np.array([1] * len(case) + [0] * len(ntrd), dtype=float)
    edge_values = np.vstack([case_edge, ntrd_edge])
    edge_rows = []
    for index, (roi_1, roi_2) in enumerate(CROSS_EDGE_NAMES):
        frame = pd.DataFrame({
            "subject_key": all_keys, "group": group, "y": edge_values[:, index]
        }).join(covariate_table, on="subject_key")
        result = fit_group_glm(frame, "y", glm_covariates)
        edge_rows.append({
            "ROI_1": roi_1,
            "ROI_2": roi_2,
            "TRD1_deviation": float(frame.loc[frame["group"] == 1, "y"].mean()),
            "nTRD_deviation": float(frame.loc[frame["group"] == 0, "y"].mean()),
            "t_group": result["t"],
            "effect_size_partial_r": result["partial_r"],
            "p_group_uncorrected": result["p"],
            "n_used": result["n"],
        })
    edge_frame = pd.DataFrame(edge_rows)
    edge_frame["q_fdr_BH_168_edges_audit_only"] = bh_fdr(edge_frame["p_group_uncorrected"])
    edge_all, edge_reported = save_group_family(
        edge_frame,
        "p_group_uncorrected",
        "q_fdr_BH_168_edges_audit_only",
        output_dir,
        "edge",
    )

    ntrd_design = sm.add_constant(ntrd[["age", "sex"]], has_constant="add").reindex(
        columns=context["regional_hc_design_columns"]
    )
    ntrd_regional = pd.DataFrame(
        (ntrd[ROI_COLUMNS].to_numpy(float) - ntrd_design.to_numpy(float) @ context["regional_beta"])
        / context["reference_sd"].to_numpy(float),
        columns=ROI_COLUMNS,
        index=ntrd.index,
    )
    ntrd_subjects = ntrd_regional.copy()
    ntrd_subjects.insert(0, "subject_key", ntrd["subject_key"].to_numpy())
    ntrd_subjects.to_excel(output_dir / "regional_nTRD_subjects.xlsx", index=False)

    case_regional = context["case_regional"]
    regional_values = np.vstack([case_regional.to_numpy(float), ntrd_regional.to_numpy(float)])
    regional_rows = []
    for index, roi in enumerate(ROI_COLUMNS):
        frame = pd.DataFrame({
            "subject_key": all_keys, "group": group, "y": regional_values[:, index]
        }).join(covariate_table, on="subject_key")
        result = fit_group_glm(frame, "y", glm_covariates)
        regional_rows.append({
            "ROI": roi,
            "TRD1_deviation": float(frame.loc[frame["group"] == 1, "y"].mean()),
            "nTRD_deviation": float(frame.loc[frame["group"] == 0, "y"].mean()),
            "t_group": result["t"],
            "effect_size_partial_r": result["partial_r"],
            "p_group_uncorrected": result["p"],
            "n_used": result["n"],
        })
    regional_frame = pd.DataFrame(regional_rows)
    regional_frame["q_fdr_BH_26_regions_audit_only"] = bh_fdr(
        regional_frame["p_group_uncorrected"]
    )
    regional_all, regional_reported = save_group_family(
        regional_frame,
        "p_group_uncorrected",
        "q_fdr_BH_26_regions_audit_only",
        output_dir,
        "regional",
    )

    context.update({
        "ntrd": ntrd,
        "ntrd_edge": ntrd_edge,
        "ntrd_regional": ntrd_regional,
        "edge_all": edge_all,
        "edge_reported": edge_reported,
        "regional_all": regional_all,
        "regional_reported": regional_reported,
    })
    return context
